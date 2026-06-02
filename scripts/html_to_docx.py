"""
Convert RAPPORT_PFE_KOFERT.html → RAPPORT_PFE_KOFERT.docx
"""
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re, os

HTML_PATH = os.path.join(os.path.dirname(__file__), '..', 'RAPPORT_PFE_KOFERT.html')
OUT_PATH  = os.path.join(os.path.dirname(__file__), '..', 'RAPPORT_PFE_KOFERT.docx')

# ── Colour palette (matches the HTML CSS) ─────────────────────────────────────
TEAL    = RGBColor(0x4E, 0xCD, 0xC4)
ORANGE  = RGBColor(0xFF, 0x6B, 0x35)
DARK    = RGBColor(0x1A, 0x1A, 0x2E)
MID     = RGBColor(0x16, 0x21, 0x3E)
TEXT    = RGBColor(0xE8, 0xE8, 0xE8)
SEC     = RGBColor(0xB0, 0xB0, 0xC0)
CODE_BG = RGBColor(0x0D, 0x0D, 0x1A)
NOTE_BG = RGBColor(0x1A, 0x2A, 0x3A)


def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_page_border(doc):
    """Light decorative border on every page via sectPr."""
    section = doc.sections[0]
    sectPr = section._sectPr
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for side in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '24')
        border.set(qn('w:color'), '4ECDC4')
        pgBorders.append(border)
    sectPr.append(pgBorders)


def style_doc(doc: Document):
    """Set default font and page margins."""
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0x20, 0x20, 0x30)
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)


def add_heading(doc: Document, text: str, level: int):
    clean = re.sub(r'\s*\[.*?\]\s*', '', text).strip()  # strip badge text
    clean = re.sub(r'\s+', ' ', clean)
    p = doc.add_heading(clean, level=level)
    run = p.runs[0] if p.runs else p.add_run(clean)
    if level == 1:
        run.font.color.rgb = TEAL
        run.font.size = Pt(16)
        run.bold = True
    elif level == 2:
        run.font.color.rgb = ORANGE
        run.font.size = Pt(13)
        run.bold = True
    else:
        run.font.color.rgb = RGBColor(0x4E, 0x90, 0xC4)
        run.font.size = Pt(11.5)
        run.bold = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after  = Pt(4)


def add_paragraph(doc: Document, text: str, bold=False, italic=False, color=None):
    text = re.sub(r'\s+', ' ', text).strip()
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_code_block(doc: Document, code: str):
    """Monospaced block with light grey shading."""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after  = Pt(4)
    para.paragraph_format.left_indent  = Cm(0.8)
    # light grey background via paragraph shading
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F0F5')
    pPr.append(shd)
    run = para.add_run(code.strip())
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x10, 0x10, 0x30)


def add_note_box(doc: Document, text: str):
    """Callout box for <div class='note'>."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent  = Cm(0.8)
    para.paragraph_format.right_indent = Cm(0.8)
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after  = Pt(6)
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E8F4F8')
    pPr.append(shd)
    run = para.add_run('ℹ  ' + re.sub(r'\s+', ' ', text).strip())
    run.font.color.rgb = RGBColor(0x1A, 0x5A, 0x7A)
    run.font.size = Pt(9.5)
    run.italic = True


def add_html_table(doc: Document, table_tag):
    rows = table_tag.find_all('tr')
    if not rows:
        return
    cols = max(len(r.find_all(['th', 'td'])) for r in rows)
    if cols == 0:
        return
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    for ri, row in enumerate(rows):
        cells = row.find_all(['th', 'td'])
        for ci, cell in enumerate(cells):
            if ci >= cols:
                break
            tc = tbl.rows[ri].cells[ci]
            text = re.sub(r'\s+', ' ', cell.get_text()).strip()
            para = tc.paragraphs[0]
            run  = para.add_run(text)
            run.font.size = Pt(9)
            if cell.name == 'th' or ri == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_bg(tc, '2C3E6B')
            else:
                bg = 'F7F9FC' if ri % 2 == 0 else 'FFFFFF'
                set_cell_bg(tc, bg)
                run.font.color.rgb = RGBColor(0x20, 0x20, 0x30)
    doc.add_paragraph()  # spacing after table


def process_inline(para, el):
    """Add inline text from a tag, preserving bold/italic/code."""
    from bs4 import NavigableString, Tag
    children = el.children if hasattr(el, 'children') else []
    for child in children:
        if isinstance(child, NavigableString):
            text = re.sub(r'\s+', ' ', str(child))
            if text.strip():
                run = para.add_run(text)
                run.font.size = Pt(10.5)
        elif isinstance(child, Tag):
            tag = child.name
            text = re.sub(r'\s+', ' ', child.get_text())
            if tag in ('strong', 'b'):
                run = para.add_run(text)
                run.bold = True
                run.font.size = Pt(10.5)
            elif tag in ('em', 'i'):
                run = para.add_run(text)
                run.italic = True
                run.font.size = Pt(10.5)
            elif tag == 'code':
                run = para.add_run(text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xC0, 0x40, 0x40)
            elif tag == 'a':
                run = para.add_run(text)
                run.font.color.rgb = TEAL
                run.font.size = Pt(10.5)
            elif tag in ('span', 'sup', 'sub', 'small'):
                process_inline(para, child)
            else:
                process_inline(para, child)


def process_list(doc: Document, ul_tag, depth=0):
    for li in ul_tag.find_all('li', recursive=False):
        # Check for nested lists inside this li
        nested = li.find(['ul', 'ol'])
        # Get direct text (excluding nested list)
        direct_text = ''
        for child in li.children:
            if hasattr(child, 'name') and child.name in ('ul', 'ol'):
                break
            if hasattr(child, 'get_text'):
                direct_text += child.get_text()
            else:
                direct_text += str(child)
        direct_text = re.sub(r'\s+', ' ', direct_text).strip()
        if direct_text:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(0.8 + depth * 0.5)
            p.paragraph_format.space_after = Pt(2)
            # Re-parse to get bold/code inline
            p.clear()
            process_inline(p, li)
        if nested:
            process_list(doc, nested, depth + 1)


def build_docx(html_path: str, out_path: str):
    with open(html_path, encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'lxml')

    doc = Document()
    style_doc(doc)

    # ── Cover page ─────────────────────────────────────────────────────────────
    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_title.paragraph_format.space_before = Pt(60)
    run = cover_title.add_run("RAPPORT DE PROJET DE FIN D\u2019\u00c9TUDES")
    run.font.size  = Pt(22)
    run.font.bold  = True
    run.font.color.rgb = TEAL

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = subtitle.add_run('Système de Monitoring Énergétique Industriel\nKOFERT — OCP Group')
    r2.font.size  = Pt(15)
    r2.font.color.rgb = ORANGE

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = meta.add_run('Application Flutter Web · Firebase · ESP32\nhttps://ocp-energy-monitor.web.app')
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(0x50, 0x50, 0x70)

    doc.add_page_break()

    # ── Parse body ─────────────────────────────────────────────────────────────
    body = soup.find('body') or soup

    skip_tags = {'script', 'style', 'head', 'nav', 'header', 'footer'}

    def walk(node, in_note=False):
        if not hasattr(node, 'name') or node.name is None:
            return
        tag = node.name

        if tag in skip_tags:
            return

        # headings
        if tag == 'h1':
            add_heading(doc, node.get_text(), 1)
            return
        if tag == 'h2':
            add_heading(doc, node.get_text(), 2)
            return
        if tag == 'h3':
            add_heading(doc, node.get_text(), 3)
            return

        # page break
        if 'page-break' in node.get('class', []):
            doc.add_page_break()
            return

        # note box
        if tag == 'div' and 'note' in node.get('class', []):
            add_note_box(doc, node.get_text())
            return

        # table
        if tag == 'table':
            add_html_table(doc, node)
            return

        # pre / code block
        if tag == 'pre':
            add_code_block(doc, node.get_text())
            return

        # paragraph
        if tag == 'p':
            text = re.sub(r'\s+', ' ', node.get_text()).strip()
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(5)
                process_inline(p, node)
            return

        # lists
        if tag in ('ul', 'ol'):
            process_list(doc, node)
            return

        # recurse into div / section / article / main
        for child in node.children:
            if hasattr(child, 'name'):
                walk(child)

    walk(body)

    doc.save(out_path)
    print(f'✓ Saved: {out_path}')


if __name__ == '__main__':
    build_docx(HTML_PATH, OUT_PATH)
